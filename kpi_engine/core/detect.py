"""Report-type detection.

Inspects an uploaded QBO export (CSV/XLSX) and determines which report it is,
along with the company name, reporting period / as-of date, and where the
column-header row sits.

QBO export layout (all reports):
  row 1: company name
  row 2: report title ("General Ledger", "A/R Aging Detail", ...)
  row 3: period — "June, 2025-May, 2026" for ranged reports,
         "As of May 31, 2026" for snapshots
  row 4: blank
  row 5: column headers
"""

from __future__ import annotations

import calendar
import csv
import datetime as dt
import re
from dataclasses import dataclass
from pathlib import Path

TB_UNSUPPORTED_MESSAGE = (
    "Trial Balance isn't needed — the General Ledger and "
    "Chart of Accounts cover it."
)

# Checked in order against the lowercased title row; first match wins.
TITLE_PATTERNS = [
    ("general ledger", "GL"),
    ("trial balance", "TB"),
    ("chart of accounts", "COA"),
    ("account list", "COA"),
    ("a/r aging detail", "AR_AGING"),
    ("a/p aging detail", "AP_AGING"),
    ("invoices and received payments", "INVOICES_PAYMENTS"),
    ("bills and applied payments", "BILLS_PAYMENTS"),
]

_MONTHS = {name.lower(): i for i, name in enumerate(calendar.month_name) if name}
_MONTHS.update({name.lower(): i for i, name in enumerate(calendar.month_abbr) if name})

# "June, 2025-May, 2026" or "June 2025 - May 2026" (comma optional, spaces loose)
_RANGE_RE = re.compile(
    r"([A-Za-z]+)\s*,?\s*(\d{4})\s*-\s*([A-Za-z]+)\s*,?\s*(\d{4})"
)
_AS_OF_RE = re.compile(r"as of\s+([A-Za-z]+)\s+(\d{1,2})\s*,?\s*(\d{4})", re.I)


@dataclass
class DetectionResult:
    report_type: str
    company_name: str | None = None
    period_start: str | None = None
    period_end: str | None = None
    as_of_date: str | None = None
    header_row_index: int | None = None
    title_text: str | None = None
    message: str | None = None


def _cell_to_str(value) -> str:
    """Convert an xlsx cell value to the string the CSV reader would produce."""
    if value is None:
        return ""
    if isinstance(value, dt.datetime):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, dt.date):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _trim_trailing(cells: list[str]) -> list[str]:
    """Drop trailing empty cells so CSV and XLSX rows compare identically
    (openpyxl pads every row to the sheet's widest column)."""
    end = len(cells)
    while end > 0 and not cells[end - 1].strip():
        end -= 1
    return cells[:end]


def read_rows(filepath: str | Path, limit: int | None = None) -> list[list[str]]:
    """Read a CSV or XLSX file into a list of string rows (one code path).

    CSV is opened with utf-8-sig (QBO exports may carry a BOM); XLSX is read
    via openpyxl and converted into the same row-of-strings structure.
    """
    path = Path(filepath)
    rows: list[list[str]] = []
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            for i, row in enumerate(wb.active.iter_rows(values_only=True)):
                if limit is not None and i >= limit:
                    break
                rows.append(_trim_trailing([_cell_to_str(v) for v in row]))
        finally:
            wb.close()
        return rows
    with open(path, encoding="utf-8-sig", newline="") as f:
        for i, row in enumerate(csv.reader(f)):
            if limit is not None and i >= limit:
                break
            rows.append(_trim_trailing([c if c is not None else "" for c in row]))
    return rows


def _row_text(rows: list[list[str]], index: int) -> str:
    """Join a row's non-empty cells back into one string.

    Unquoted commas in QBO's title/period/footer lines can split a single
    logical value across CSV cells; rejoining restores it.
    """
    if index >= len(rows):
        return ""
    return ",".join(c for c in rows[index] if c.strip()).strip()


def _month_bounds(month_name: str, year: str) -> tuple[dt.date, dt.date] | None:
    month = _MONTHS.get(month_name.lower())
    if month is None:
        return None
    y = int(year)
    last = calendar.monthrange(y, month)[1]
    return dt.date(y, month, 1), dt.date(y, month, last)


def _parse_period(text: str) -> tuple[str | None, str | None, str | None]:
    """Return (period_start, period_end, as_of_date) as ISO strings."""
    m = _AS_OF_RE.search(text)
    if m:
        month = _MONTHS.get(m.group(1).lower())
        if month is not None:
            try:
                d = dt.date(int(m.group(3)), month, int(m.group(2)))
            except ValueError:
                return None, None, None
            return None, None, d.isoformat()
    m = _RANGE_RE.search(text)
    if m:
        start = _month_bounds(m.group(1), m.group(2))
        end = _month_bounds(m.group(3), m.group(4))
        if start and end:
            return start[0].isoformat(), end[1].isoformat(), None
    return None, None, None


def _find_header_row(rows: list[list[str]], start: int = 3) -> int | None:
    """First non-blank row at/after `start` (row 4 is blank, row 5 is headers)."""
    for i in range(start, len(rows)):
        if any(c.strip() for c in rows[i]):
            return i
    return None


def detect_from_rows(rows: list[list[str]]) -> DetectionResult:
    """Detect report type from already-read rows (first ~10 are enough)."""
    company = _row_text(rows, 0) or None
    title = _row_text(rows, 1)
    period_start, period_end, as_of = _parse_period(_row_text(rows, 2))
    header_idx = _find_header_row(rows)

    report_type = "UNKNOWN"
    title_lower = title.lower()
    for pattern, rtype in TITLE_PATTERNS:
        if pattern in title_lower:
            report_type = rtype
            break

    message = TB_UNSUPPORTED_MESSAGE if report_type == "TB" else None
    return DetectionResult(
        report_type=report_type,
        company_name=company,
        period_start=period_start,
        period_end=period_end,
        as_of_date=as_of,
        header_row_index=header_idx,
        title_text=title,
        message=message,
    )


def detect_report_type(filepath: str | Path) -> DetectionResult:
    """Identify which QBO report a CSV/XLSX export is, plus its period."""
    return detect_from_rows(read_rows(filepath, limit=10))
